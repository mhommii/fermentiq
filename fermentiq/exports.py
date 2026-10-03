"""Downloadable outputs, all built in memory: PDF report, Excel workbook, chart images (ZIP)."""

import io
import zipfile
from datetime import datetime
from pathlib import Path

import matplotlib
import pandas as pd
from fpdf import FPDF
from fpdf.fonts import FontFace
from openpyxl.styles import Font

from . import about
from .charts import ALL_CHARTS, render_png
from .html_report import ISSUE_NAMES, SEVERITY_LABELS, run_title, sample_label
from .report import interpretation, questions_for_review
from .style import get_theme

FONT_DIR = Path(matplotlib.get_data_path()) / "fonts" / "ttf"  # DejaVu Sans ships with matplotlib (has μ, ⁻¹, →)


def _rgb(hex_color):
    hex_color = hex_color.lstrip("#")
    return tuple(int(hex_color[i:i + 2], 16) for i in (0, 2, 4))


# ---------- PDF ----------

class _ReportPDF(FPDF):
    def __init__(self, theme):
        super().__init__(format="A4")
        self.t = theme
        self.add_font("DejaVu", "", str(FONT_DIR / "DejaVuSans.ttf"))
        self.add_font("DejaVu", "B", str(FONT_DIR / "DejaVuSans-Bold.ttf"))
        self.set_margins(14, 14, 14)
        self.set_auto_page_break(True, margin=16)
        self.set_draw_color(*_rgb(theme.grid))

    def header(self):
        self.set_fill_color(*_rgb(self.t.surface))  # page background in the theme color
        self.rect(0, 0, self.w, self.h, "F")

    def footer(self):
        self.set_y(-11)
        self.font(7.5, color=self.t.muted)
        self.cell(0, 4, f"{about.NAME} · built by {about.AUTHOR} · {about.GITHUB_URL}", align="L")
        self.cell(0, 4, f"Page {self.page_no()}/{{nb}}", align="R")

    def font(self, size, bold=False, color=None):
        self.set_font("DejaVu", "B" if bold else "", size)
        self.set_text_color(*_rgb(color or self.t.ink))

    def section(self, number, title):
        self.ln(4)
        self.font(12, bold=True, color=self.t.accent)
        self.cell(7, 7, str(number))
        self.font(12, bold=True)
        self.cell(0, 7, title, new_x="LMARGIN", new_y="NEXT")
        self.ln(1)

    def note(self, text):
        self.font(8, color=self.t.muted)
        self.multi_cell(0, 4, text, new_x="LMARGIN", new_y="NEXT")
        self.ln(1)

    def table_rows(self, headers, rows, widths, size=8):
        self.font(size, color=self.t.ink_secondary)
        surface = _rgb(self.t.surface)
        self.set_fill_color(*surface)  # tables have no colored fill, whatever color was used last
        heading = FontFace(emphasis="BOLD", color=_rgb(self.t.ink_secondary), fill_color=surface)
        with self.table(col_widths=widths, text_align="LEFT", borders_layout="HORIZONTAL_LINES",
                        line_height=size * 0.55, padding=(1.2, 1.5), headings_style=heading,
                        cell_fill_mode="NONE", first_row_as_headings=True) as table:
            for values in [headers] + rows:
                row = table.row()
                for value in values:
                    row.cell(str(value))
        self.ln(2)


def build_pdf(record, findings, kin, theme=None):
    """The report as PDF bytes, in the chosen theme."""
    t = get_theme(theme)
    pdf = _ReportPDF(t)
    pdf.add_page()
    df, mu, fit = kin["table"], kin["mu_max"], kin["logistic"]
    organism, equipment, group, run_date = run_title(record, kin)

    # header
    pdf.font(17, bold=True)
    pdf.cell(0, 9, "FermentIQ · Batch Record Review", new_x="LMARGIN", new_y="NEXT")
    pdf.font(9.5, color=t.ink_secondary)
    pdf.cell(0, 5, f"{organism} growth in {equipment} · {group} · {run_date} · {len(df)} samples · "
                   f"{df['t_min'].max():.0f} min run", new_x="LMARGIN", new_y="NEXT")
    pdf.font(7.5, color=t.muted)
    pdf.cell(0, 5, f"Source: {Path(record.source).name} · Generated {datetime.now():%d %b %Y, %H:%M} · "
                   f"{t.label} theme", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)

    # summary cards
    counts = findings["severity"].value_counts()
    cards = [("Max specific growth rate", f"{mu['mu_max_per_h']:.2f} h⁻¹", f"R² {mu['r2']:.3f}"),
             ("Doubling time", f"{mu['doubling_time_min']:.0f} min", "ln 2 / μmax"),
             ("Maximum OD600", f"{fit['K_od']:.2f}", f"Logistic fit · R² {fit['r2']:.3f}"),
             ("Findings", " / ".join(str(counts.get(k, 0)) for k in SEVERITY_LABELS), "critical / major / minor")]
    width, top = (pdf.w - 28 - 9) / 4, pdf.get_y()
    for i, (label, value, sub) in enumerate(cards):
        x = 14 + i * (width + 3)
        pdf.set_draw_color(*_rgb(t.baseline))
        pdf.rect(x, top, width, 21)
        pdf.set_xy(x + 3, top + 2.5)
        pdf.font(7.5, color=t.ink_secondary)
        pdf.cell(width - 6, 4, label)
        pdf.set_xy(x + 3, top + 7.5)
        pdf.font(14, bold=True)
        pdf.cell(width - 6, 7, value)
        pdf.set_xy(x + 3, top + 15)
        pdf.font(7, color=t.muted)
        pdf.cell(width - 6, 4, sub)
    pdf.set_draw_color(*_rgb(t.grid))
    pdf.set_y(top + 25)

    # 1 findings
    pdf.section(1, "Data integrity findings")
    pdf.note("Every calculation and entry in the batch record was re-checked. Recorded values are shown next "
             "to what they should be.")
    for key, label in SEVERITY_LABELS.items():
        rows = findings[findings["severity"] == key]
        if not len(rows):
            continue
        pdf.ln(1)
        if pdf.will_page_break(25):  # don't leave a severity heading alone at the bottom of a page
            pdf.add_page()
        pdf.set_fill_color(*_rgb(t.status[key]))
        pdf.ellipse(pdf.l_margin, pdf.get_y() + 1.2, 2.8, 2.8, "F")
        pdf.set_x(pdf.l_margin + 4.5)
        pdf.font(9.5, bold=True)
        pdf.cell(0, 5, f"{label} ({len(rows)})", new_x="LMARGIN", new_y="NEXT")
        pdf.table_rows(["Sample", "Issue", "Recorded → Expected", "Why it matters"],
                       [[sample_label(r["sample"]), ISSUE_NAMES.get(r["check"], r["check"]),
                         f"{r['recorded']} → {r['expected']}", r["message"]] for _, r in rows.iterrows()],
                       widths=(17, 35, 50, 80))

    # 2 kinetics
    pdf.add_page()
    pdf.section(2, "Growth kinetics")
    excluded = ", ".join(f"S{s}" for s in fit["excluded_samples"])
    pdf.note("Time from clock time. OD600 uses the dilution-corrected value where a sample was diluted."
             + (f" Excluded from fit: {excluded} (read near instrument limit)." if excluded else ""))
    for name in ("growth_curve", "specific_growth_rate"):
        pdf.image(io.BytesIO(render_png(ALL_CHARTS[name], kin, t)), w=pdf.epw)
        pdf.ln(2)
    pdf.font(9.5, bold=True)
    pdf.cell(0, 6, "Key observations", new_x="LMARGIN", new_y="NEXT")
    pdf.font(8.5, color=t.ink_secondary)
    for line in interpretation(findings, kin).splitlines():
        pdf.multi_cell(0, 4.3, "•  " + line.lstrip("- "), new_x="LMARGIN", new_y="NEXT")
        pdf.ln(0.8)
    pdf.ln(2)
    if pdf.will_page_break(12 + 6 * (len(df) + 1)):  # keep the short sample table on one page
        pdf.add_page()
    pdf.font(9.5, bold=True)
    pdf.cell(0, 6, "Sample data", new_x="LMARGIN", new_y="NEXT")
    pdf.table_rows(["Sample", "Time (min)", "OD600", "pH probe", "pH meter", "Wet pellet (mg)", "Note"],
                   [[f"S{r['sample']}", f"{r['t_min']:.0f}", f"{r['od600']:.3f}", f"{r['ph_probe']:.2f}",
                     f"{r['ph_meter']:.2f}",
                     "–" if pd.isna(r["pellet_recomputed_g"]) else f"{r['pellet_recomputed_g'] * 1000:.1f}",
                     r["od_flag"].capitalize()] for _, r in df.iterrows()],
                   widths=(14, 20, 18, 18, 18, 26, 40))

    # 3 process (images move to a new page by themselves if they don't fit)
    pdf.section(3, "Process and instruments")
    for name in ("ph", "pellet_vs_od"):
        pdf.image(io.BytesIO(render_png(ALL_CHARTS[name], kin, t)), w=pdf.epw)
        pdf.ln(2)

    # 4 questions
    pdf.section(4, "Questions for review")
    pdf.font(8.5, color=t.ink_secondary)
    for line in questions_for_review(findings, df).splitlines():
        pdf.multi_cell(0, 4.3, line, new_x="LMARGIN", new_y="NEXT")
        pdf.ln(0.8)
    return bytes(pdf.output())


# ---------- Excel ----------

def build_excel(record, findings, kin):
    """The data behind the report as an Excel workbook (bytes)."""
    df, mu, fit, settings = kin["table"], kin["mu_max"], kin["logistic"], kin["settings"]
    organism, equipment, group, run_date = run_title(record, kin)
    counts = findings["severity"].value_counts()
    summary = pd.DataFrame([
        ("Source file", Path(record.source).name), ("Organism", organism), ("Equipment", equipment),
        ("Group", group), ("Run date", run_date), ("Samples", len(df)),
        ("Run length (min)", round(df["t_min"].max())),
        ("μmax (1/h)", round(mu["mu_max_per_h"], 4)), ("μmax R²", round(mu["r2"], 4)),
        ("μmax window", f"S{mu['first_sample']}–S{mu['last_sample']}"),
        ("Doubling time (min)", round(mu["doubling_time_min"], 1)),
        ("Max OD600, logistic K", round(fit["K_od"], 4)), ("Logistic fit R²", round(fit["r2"], 4)),
        ("Excluded from fit", ", ".join(f"S{s}" for s in fit["excluded_samples"]) or "none"),
        ("Critical findings", counts.get("critical", 0)), ("Major findings", counts.get("major", 0)),
        ("Minor findings", counts.get("minor", 0)), ("Generated", f"{datetime.now():%Y-%m-%d %H:%M}"),
    ], columns=["Item", "Value"])
    findings_sheet = pd.DataFrame({
        "Severity": findings["severity"].astype(str).str.capitalize(),
        "Issue": findings["check"].map(lambda c: ISSUE_NAMES.get(c, c)),
        "Sample": findings["sample"].map(sample_label),
        "Field": findings["field"], "Recorded": findings["recorded"], "Expected": findings["expected"],
        "Why it matters": findings["message"],
    })
    samples = df[["sample", "t_min", "time_point_min", "od600", "od_flag", "ph_probe", "ph_meter",
                  "pellet_recomputed_g", "wet_cells_g_per_L"]].rename(columns={
        "sample": "Sample", "t_min": "Time from clock (min)", "time_point_min": "Recorded time point (min)",
        "od600": "OD600 (corrected)", "od_flag": "OD note", "ph_probe": "pH probe", "ph_meter": "pH meter",
        "pellet_recomputed_g": "Wet pellet, recomputed (g)", "wet_cells_g_per_L": "Wet cells (g/L)"})
    rates = kin["rates"].rename(columns={
        "from_sample": "From sample", "to_sample": "To sample", "t_start_h": "Start (h)", "t_end_h": "End (h)",
        "mu_per_h": "μ (1/h)", "ph_probe_end": "pH probe at end", "phase": "Phase"})
    raw = lambda frame: frame[frame["used"]].drop(columns=["used"]) if len(frame) else frame
    settings_sheet = pd.DataFrame(settings.as_rows(), columns=["Setting", "Value"])

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        for name, frame in [("Summary", summary), ("Findings", findings_sheet), ("Samples", samples),
                            ("Growth rates", rates), ("Process (raw)", raw(record.process)),
                            ("Pellets (raw)", raw(record.pellets)), ("OD600 (raw)", raw(record.od)),
                            ("Signatures", record.signatures), ("Settings", settings_sheet)]:
            frame.to_excel(writer, sheet_name=name, index=False)
            sheet = writer.sheets[name]
            sheet.freeze_panes = "A2"
            for cell in sheet[1]:
                cell.font = Font(bold=True)
            for column in sheet.columns:  # rough auto-width, capped so long text stays readable
                longest = max(len(str(c.value)) if c.value is not None else 0 for c in column)
                sheet.column_dimensions[column[0].column_letter].width = min(max(10, longest + 2), 60)
    return buffer.getvalue()


# ---------- chart images ----------

def build_chart_zip(kin, theme=None):
    """All four charts as PNG files in one ZIP (bytes)."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, plot_fn in ALL_CHARTS.items():
            archive.writestr(f"{name}.png", render_png(plot_fn, kin, theme))
    return buffer.getvalue()


def write_all(record, findings, kin, report_dir, theme=None):
    """CLI helper: write the PDF, Excel workbook and chart PNGs next to report.html."""
    report_dir = Path(report_dir)
    (report_dir / "charts").mkdir(parents=True, exist_ok=True)
    (report_dir / "report.pdf").write_bytes(build_pdf(record, findings, kin, theme))
    (report_dir / "data.xlsx").write_bytes(build_excel(record, findings, kin))
    for name, plot_fn in ALL_CHARTS.items():
        (report_dir / "charts" / f"{name}.png").write_bytes(render_png(plot_fn, kin, theme))
